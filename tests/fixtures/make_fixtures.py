"""Regenerates the import fixtures used by the tests (t9, t21).

    pip install python-docx reportlab
    python3 tests/fixtures/make_fixtures.py

test.docx     title, paragraphs with bold/italic, a heading with bullets, a subheading with a numbered list, a table
test.pdf      two pages: big title, a paragraph wrapped over three lines, a subheading, bullets, a numbered list,
              and a footer repeated on every page (the importer drops repeated headers and footers)
textbook.pdf  three two-column pages: running header with page numbers, numbered section headings, long wrapped
              paragraphs, a sidebar heading, bullets and a figure caption
"""
import os

import docx
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

HERE = os.path.dirname(os.path.abspath(__file__))


def make_docx(path):
    d = docx.Document()
    d.add_paragraph('Community Garden Operations Manual', style='Title')
    d.add_paragraph('Intro paragraph with some bold text.')
    p = d.add_paragraph('Mixed ')
    p.add_run('bold').bold = True
    p.add_run(' and ')
    p.add_run('italic').italic = True
    d.add_paragraph('Scope', style='Heading 1')
    d.add_paragraph('First item', style='List Bullet')
    d.add_paragraph('Second item', style='List Bullet')
    d.add_paragraph('Steps', style='Heading 2')
    d.add_paragraph('Do this', style='List Number')
    d.add_paragraph('Then that', style='List Number')
    t = d.add_table(rows=2, cols=2)
    for r, row in enumerate([['Col A', 'Col B'], ['1', '2']]):
        for c, text in enumerate(row):
            t.cell(r, c).text = text
    cp = d.core_properties
    cp.author = cp.last_modified_by = ''
    cp.comments = cp.title = ''
    d.save(path)


def new_canvas(path):
    c = canvas.Canvas(path, pagesize=letter, invariant=1)
    c.setAuthor('')
    c.setCreator('')
    c.setTitle('')
    return c


def make_test_pdf(path):
    c = new_canvas(path)
    para = ['This document describes the garden layout and the way',
            'volunteers are assigned to plots across seasons. Paragraphs are',
            'wrapped across several lines in the PDF.']
    for n, title in enumerate(['Community Garden Overview', 'Seasons'], 1):
        c.setFont('Helvetica-Bold', 22)
        c.drawString(72, 720, title)
        c.setFont('Helvetica', 11)
        for i, line in enumerate(para):
            c.drawString(72, 690 - 15 * i, line)
        c.setFont('Helvetica-Bold', 15)
        c.drawString(72, 615, 'Key points')
        c.setFont('Helvetica', 11)
        for i, line in enumerate(['• First bullet point here', '• Second bullet point here', '1. Numbered step one', '2. Numbered step two']):
            c.drawString(72, 593 - 15 * i, line)
        c.setFont('Helvetica', 9)
        c.drawString(300, 40, str(n))
        c.drawString(72, 40, 'Riverside Garden Club — Internal')
        c.showPage()
    c.save()


BODY = ['Photosynthesis converts light energy into', 'chemical energy stored in glucose. The', 'process takes place in the chloroplasts,',
        'where chlorophyll absorbs light most', 'efficiently in the blue and red']
BODY += ['wavelengths.Photosynthesis converts light', 'energy into chemical energy stored in', 'glucose. The process takes place in the',
         'chloroplasts, where chlorophyll absorbs', 'light most efficiently in the blue and red'] * 2
BODY += ['wavelengths.']


def make_textbook_pdf(path):
    c = new_canvas(path)
    body = lambda x, y, lines: [c.drawString(x, y - 13 * i, line) for i, line in enumerate(lines)] and y - 13 * (len(lines) - 1)
    for page in range(3):
        c.setFont('Helvetica', 8)
        c.drawString(72, 752, 'CHAPTER 4  ·  ENERGY IN LIVING SYSTEMS')
        c.drawString(526.656, 752, str(112 + page))
        # left column: (section heading) paragraph, numbered subheading, a paragraph that runs on
        y = 712
        if page == 0:
            c.setFont('Helvetica-Bold', 11)
            c.drawString(72, 712, '4.2 Light-Dependent Reactions')
            y = 690
        c.setFont('Times-Roman', 10.5)
        y = body(72, y, BODY) - 23
        c.setFont('Helvetica-Bold', 10.5)
        c.drawString(72, y, '4.2.%d The Thylakoid Membrane' % (page + 1))
        c.setFont('Times-Roman', 10.5)
        body(72, y - 16, BODY[:10])
        # right column: (sidebar heading) paragraph, bullets, figure caption
        y = 690 if page == 0 else 712
        if page == 1:
            c.setFont('Helvetica-Bold', 10.5)
            c.drawString(316, 712, 'KEY CONCEPTS')
            y = 694
        c.setFont('Times-Roman', 10.5)
        y = body(316, y, BODY) - 23
        y = body(316, y, ['• Light is absorbed by Photosystem II.', '• Electrons move along the chain.', '• ATP synthase produces ATP.']) - 19
        c.setFont('Times-Italic', 8.5)
        c.drawString(316, y, 'Figure 4.3 The thylakoid membrane in cross-section.')
        c.showPage()
    c.save()


if __name__ == '__main__':
    make_docx(os.path.join(HERE, 'test.docx'))
    make_test_pdf(os.path.join(HERE, 'test.pdf'))
    make_textbook_pdf(os.path.join(HERE, 'textbook.pdf'))
    print('wrote test.docx, test.pdf, textbook.pdf in', HERE)
